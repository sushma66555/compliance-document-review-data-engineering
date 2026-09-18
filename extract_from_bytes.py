import io
from docx import Document
import pdfplumber
import openpyxl


def extract_text_from_bytes(file_bytes, content_type):
    """
    Detects file type from content_type and extracts text accordingly.
    Works directly on bytes, no need to save the file to disk first.
    """
    file_stream = io.BytesIO(file_bytes)

    if "pdf" in content_type:
        return _extract_pdf(file_stream)
    elif "word" in content_type or "docx" in content_type or "officedocument.wordprocessingml" in content_type:
        return _extract_docx(file_stream)
    elif "excel" in content_type or "spreadsheet" in content_type:
        return _extract_xlsx(file_stream)
    else:
        raise ValueError(f"Unsupported content type: {content_type}")


def _extract_docx(file_stream):
    doc = Document(file_stream)
    text_parts = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(text_parts)


def _extract_pdf(file_stream):
    text_parts = []
    with pdfplumber.open(file_stream) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def _extract_xlsx(file_stream):
    wb = openpyxl.load_workbook(file_stream)
    sheet = wb.active
    text_parts = []
    for row in sheet.iter_rows(values_only=True):
        row_text = " | ".join(str(c) for c in row if c is not None)
        if row_text.strip():
            text_parts.append(row_text)
    return "\n".join(text_parts)